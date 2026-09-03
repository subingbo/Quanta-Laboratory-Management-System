from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Quanta后台管理系统接口对照说明.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "263445"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
GREEN = "EAF7EF"
AMBER = "FFF4DF"
RED = "FDECEC"
WHITE = "FFFFFF"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=9, color=INK, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade_status(cell, status):
    if status in ("已有", "可复用"):
        set_cell_fill(cell, GREEN)
    elif status in ("待确认", "需扩展"):
        set_cell_fill(cell, AMBER)
    elif status in ("缺失", "高风险"):
        set_cell_fill(cell, RED)


def add_table(doc, headers, rows, widths, font_size=8.3, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, value in enumerate(headers):
        set_cell_fill(hdr.cells[i], LIGHT_BLUE)
        p = hdr.cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(value), size=8.5, color=DARK_BLUE, bold=True)
    for values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for i, value in enumerate(values):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(str(value)), size=font_size)
            if status_col == i:
                shade_status(row.cells[i], str(value))
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_note(doc, label, text, fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_row_cant_split(table.rows[0])
    set_table_geometry(table, [9360])
    set_cell_fill(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label + "："), size=9.5, color=DARK_BLUE, bold=True)
    set_run_font(p.add_run(text), size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.add_run("Quanta 后台管理系统｜接口对照说明"), size=8.5, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("技术接口核对文档"), size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run("Quanta 后台管理系统接口对照说明"), size=24, color="172B4D", bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    set_run_font(p.add_run("后端现有接口 × 前端定义接口 × 缺失能力"), size=13, color=MUTED)

    meta = [
        ("文档用途", "供前后端逐项核对接口覆盖范围、字段契约与权限边界"),
        ("依据", "项目根目录《接口文档.md》、src/api 前端请求定义、当前管理端产品规则"),
        ("生成日期", str(date.today())),
        ("项目状态", "后端尚未上线；开发环境当前以 Mock 数据运行"),
    ]
    add_table(doc, ["项目", "说明"], meta, [1800, 7560], font_size=9)
    add_note(doc, "阅读方式", "优先查看第 3 节。红色“缺失”表示后端文档和已知 OpenAPI 均未确认；黄色“待确认/需扩展”表示已有相近接口，但路径、字段或权限仍需后端确认。", AMBER)


def main():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "1. 对照结论", 1)
    summary_rows = [
        ("后端 Markdown 已明确", "认证、个人资料、文件、成员名录、活动/报名、图书、工位、塔服、面试投递，以及 9 组标准管理 CRUD。"),
        ("前端已直接调用", "认证、成员名录、面试结果录入、预约记录、借阅记录、塔服订单记录等。"),
        ("前端已定义但后端缺失/待确认", "控制台统计、届次与留任、招新管理端完整流程、宣讲会/精英分享会聚合名单、学习资料、塔服收款确认。"),
        ("最高风险", "后台 CRUD 权限注解被注释；前端隐藏按钮不能替代后端鉴权。"),
    ]
    add_table(doc, ["类别", "结论"], summary_rows, [2100, 7260], font_size=9)

    doc.add_page_break()
    add_heading(doc, "1.1 状态口径", 2)
    add_table(doc, ["状态", "含义"], [
        ("已有", "根目录《接口文档.md》已明确记录，可进入联调；仍应核对实际返回字段。"),
        ("可复用", "已有通用接口可能覆盖页面，但需要统一查询参数、字段映射或返回结构。"),
        ("待确认", "前端已定义；Markdown 未列出，但旧 OpenAPI/Apifox 中曾发现相近接口，需后端确认。"),
        ("需扩展", "已有接口，但当前字段、数据范围、上传目录或业务事务不足。"),
        ("缺失", "当前后端文档未提供，前端只能使用 Mock。"),
    ], [1500, 7860], font_size=9, status_col=0)

    add_heading(doc, "2. 后端接口文档：现有接口及用途", 1)
    add_note(doc, "通用约定", "AjaxResult 使用 code/msg/data；分页列表使用 TableDataInfo 的 total/rows/code/msg；除免登录接口外使用 Authorization: Bearer <token>。", LIGHT_BLUE)

    add_heading(doc, "2.1 认证与个人信息", 2)
    auth_rows = [
        ("GET", "/captchaImage", "获取验证码开关、验证码图片与 uuid。"),
        ("POST", "/login", "账号密码登录，返回 token 及成员扩展身份信息。"),
        ("GET", "/getInfo", "获取当前用户、角色、权限、部门与成员身份。"),
        ("GET", "/getRouters", "获取当前用户可访问的动态菜单与前端路由树。"),
        ("POST", "/register", "新生自助注册。"),
        ("POST", "/logout", "退出登录并使当前令牌失效。"),
        ("POST", "/unlockscreen", "校验密码并解除锁屏。"),
        ("GET", "/system/user/profile", "查询当前用户个人资料。"),
        ("PUT", "/system/user/profile", "修改当前用户基础资料。"),
        ("PUT", "/system/user/profile/updatePwd", "修改当前用户密码。"),
        ("POST", "/system/user/profile/avatar", "上传并更新当前用户头像。"),
    ]
    add_table(doc, ["方法", "接口", "用途"], auth_rows, [900, 3300, 5160])

    add_heading(doc, "2.2 通用文件", 2)
    add_table(doc, ["方法", "接口", "用途/限制"], [
        ("POST", "/common/upload", "上传单个文件；文档仅允许 qt/clothing-item、qt/payment-qr、qt/payment-proof 子目录。"),
        ("GET", "/common/download", "按文件名下载服务器文件。"),
    ], [900, 3300, 5160])

    add_heading(doc, "2.3 成员、活动与报名", 2)
    add_table(doc, ["方法", "接口", "用途"], [
        ("GET", "/qt/member/list", "查询在职普通 Quanta 成员；支持姓名、部门、届次筛选。"),
        ("GET", "/system/activity/list", "分页查询活动列表。"),
        ("GET", "/system/activity/{activityId}", "查询活动详情与容量等信息。"),
        ("POST", "/system/signup", "C 端用户报名活动。"),
        ("GET", "/system/signup/detailList?userId=...", "查询报名明细联表记录；文档以用户维度为主。"),
    ], [900, 3600, 4860])

    add_heading(doc, "2.4 图书借阅", 2)
    add_table(doc, ["方法", "接口", "用途"], [
        ("GET", "/system/book/list", "分页查询图书。"),
        ("GET", "/system/book/{bookId}", "查询单本图书详情。"),
        ("POST", "/system/borrow", "创建借阅记录。"),
        ("GET", "/system/borrow/detailList?userId=...", "查询带借阅人和图书信息的借阅记录。"),
    ], [900, 3600, 4860])

    add_heading(doc, "2.5 工位预约", 2)
    add_table(doc, ["方法", "接口", "用途"], [
        ("GET", "/system/workstation/list", "分页查询工位。"),
        ("GET", "/system/workstation/{id}", "查询工位详情。"),
        ("POST", "/system/reservation", "创建工位预约。"),
        ("GET", "/system/reservation/detailList?userId=...", "查询预约记录联表数据。"),
    ], [900, 3600, 4860])

    add_heading(doc, "2.6 塔服订购", 2)
    add_table(doc, ["方法", "接口", "用途"], [
        ("GET", "/system/item/list", "查询服装款式、颜色和尺码配置。"),
        ("GET", "/system/item/{id}", "查询服装款式详情。"),
        ("GET", "/system/payment-config/list", "查询固定付款码配置。"),
        ("POST", "/system/order", "创建塔服订单。"),
        ("GET", "/system/order/detailList?userId=...", "查询订单联表明细。"),
    ], [900, 3600, 4860])

    add_heading(doc, "2.7 招新投递与面试结果", 2)
    add_table(doc, ["方法", "接口", "用途"], [
        ("POST", "/qt/interview/apply", "新生提交招新申请与两个部门志愿。"),
        ("GET", "/qt/interview/my", "查询当前用户自己的投递。"),
        ("GET", "/qt/interview/myResults", "查询当前用户自己的面试结果。"),
        ("POST", "/qt/interview/result", "管理端录入面试结果；当前文档注明没有权限控制。"),
    ], [900, 3600, 4860])

    add_heading(doc, "2.8 后台标准 CRUD 接口", 2)
    add_note(doc, "统一形态", "每个资源默认提供 GET /list、GET /{id}、POST 新增、PUT 修改、DELETE /{ids} 删除、POST /export 导出 Excel。", LIGHT_BLUE)
    crud_rows = [
        ("实验室图书", "/system/book", "bookId", "图书资料维护。"),
        ("实验室活动", "/system/activity", "activityId", "活动发布与维护。"),
        ("活动报名", "/system/signup", "signupId", "报名记录维护，另有 detailList。"),
        ("图书借阅记录", "/system/borrow", "borrowId", "借阅、归还与逾期记录维护。"),
        ("实验室工位", "/system/workstation", "workstationId", "工位配置维护。"),
        ("工位预约记录", "/system/reservation", "reservationId", "预约记录维护。"),
        ("服装配置", "/system/item", "itemId", "multipart 新增/修改；可清空效果图。"),
        ("服装订单", "/system/order", "orderId", "订单记录维护。"),
        ("付款码配置", "/system/payment-config", "configId", "multipart 新增/修改；可清空二维码。"),
        ("实验室成员", "/qt/member", "userId", "当前仅提供 list，不是完整 CRUD。"),
    ]
    add_table(doc, ["资源", "路径前缀", "主键", "用途/备注"], crud_rows, [1800, 2700, 1500, 3360], font_size=8.1)
    add_note(doc, "高风险", "《接口文档.md》明确写明上述后台 CRUD 的 @PreAuthorize 当前全部被注释，任何登录用户均可调用。必须在后端恢复接口级鉴权与部门数据权限。", RED)

    add_heading(doc, "3. 前端已定义、后端缺失或待确认的接口", 1)
    add_heading(doc, "3.1 前端 src/api 中已定义", 2)
    missing_rows = [
        ("缺失", "GET", "/dashboard/stats", "控制台统计成员数、今日简历、待处理预约、待确认收款等卡片数据。", "P1"),
        ("缺失", "GET", "/qt/member/cohorts", "返回当前届与历史届 tab，用于成员名单切换。", "P0"),
        ("缺失", "PUT", "/qt/member/{userId}/retain", "确认成员留任；更新原届状态、必要时新建下一届并同步成员，要求幂等事务。", "P0"),
        ("待确认", "DELETE", "/system/user/{userId}", "CEO 删除成员账号/成员记录；旧 OpenAPI 曾存在，Markdown 未列出。", "P0"),
        ("待确认", "PUT", "/system/user/resetPwd", "CEO 将指定成员密码重置为约定初始值；旧 OpenAPI 曾存在。", "P0"),
        ("待确认", "POST", "/system/user/importData", "按 Excel 模板批量导入成员；需返回成功/失败数量和逐行错误。", "P1"),
        ("待确认", "POST", "/system/user/importTemplate", "下载成员导入模板；应返回文件流。", "P1"),
        ("缺失", "GET", "/qt/interview/admin/statistics", "招聘看板统计各阶段人数和待处理数。", "P1"),
        ("缺失", "GET", "/qt/interview/admin/applications", "管理端按轮次、部门、状态筛选候选人；必须落实部门数据范围。", "P0"),
        ("缺失", "GET", "/qt/interview/admin/applications/{id}", "查看候选人完整简历、两个志愿、各轮状态与投递信息。", "P0"),
        ("缺失", "GET", "/qt/interview/admin/evaluations", "按申请、轮次、部门查询多位面试官的面评。", "P0"),
        ("缺失", "POST", "/qt/interview/admin/evaluations", "经理层创建或更新自己的面评。", "P0"),
        ("缺失", "POST", "/qt/interview/admin/offers", "二面 Pass/Out 后原子写入最终结果并生成录用/淘汰通知。", "P0"),
        ("缺失", "GET", "/qt/interview/admin/applications/export", "导出当前筛选结果；需确认 GET/POST、筛选参数和文件流。", "P1"),
        ("缺失", "GET", "/qt/activity/lecture/registrations", "宣讲会报名名单，返回报名人数/名额及姓名、学号、专业、电话、时间。", "P1"),
        ("缺失", "GET", "/qt/activity/sharing/registrations", "精英分享会报名名单，返回报名人数/名额及联系方式、备注。", "P1"),
        ("缺失", "GET", "/qt/materials", "查询学习资料列表及分类、大小、上传者、时间、可见性。", "P1"),
        ("缺失", "POST", "/qt/materials", "上传资料文件并保存分类、可见范围等元数据。", "P1"),
        ("缺失", "DELETE", "/qt/materials/{materialId}", "删除学习资料文件及元数据，需校验上传者/管理权限。", "P1"),
        ("缺失", "PUT", "/system/order/{orderId}/approve", "管理端确认塔服收款；校验订单状态、记录操作者与时间。", "P0"),
    ]
    add_table(doc, ["状态", "方法", "接口", "用途", "优先级"], missing_rows, [1050, 750, 2920, 4040, 600], font_size=7.7, status_col=0)

    add_heading(doc, "3.2 项目需要但前端尚未封装的后端能力", 2)
    extra_rows = [
        ("GET", "/qt/interview/admin/applications/{applicationId}/results", "查看该申请各志愿、各轮次的评定结果，避免无法判断评价的是哪个志愿。", "P0"),
        ("PUT", "/qt/interview/admin/evaluations/{evaluationId}", "修改自己的面评；需要面评归属与角色校验。", "P0"),
        ("PUT", "/qt/interview/admin/applications/{id}/join-status", "更新是否留任/是否接受等确认状态（若与 offers 分离）。", "P1"),
        ("PUT", "/qt/interview/admin/applications/{id}/final-status", "管理员修正最终状态并保留审计记录。", "P1"),
        ("GET", "/qt/materials/{materialId}/download", "学习资料下载/预览，返回文件流并校验可见范围。", "P1"),
        ("PUT", "/qt/materials/{materialId}", "修改资料分类、名称或可见范围。", "P2"),
        ("后台定时任务", "每年 8 月 1 日成员换届", "管理层自动卸任；经理/实习生未确认留任则卸任；留任者同步到下一届。不是前端接口。", "P0"),
        ("配置/字典", "招新轮次与部门字典", "避免轮次、部门、状态在前端写死；支持后续届次与组织调整。", "P2"),
    ]
    add_table(doc, ["方法/类型", "建议接口", "用途", "优先级"], extra_rows, [1300, 3300, 4160, 600], font_size=7.9)

    add_heading(doc, "4. 可复用接口与需要扩展的地方", 1)
    reuse_rows = [
        ("成员管理", "/qt/member/list", "需扩展", "补充 roleCategory、memberStatus、joinTime、canRetain；支持历史届、已卸任、角色筛选。当前只返回在职普通成员。"),
        ("宣讲会/分享会", "/system/activity/list + /system/signup/detailList", "可复用", "若 detailList 支持 activityId 且返回学号、专业、电话、联系方式、备注和 capacity，则可不新增聚合路径。"),
        ("图书借阅", "/system/borrow/detailList", "需扩展", "页面需要 bookType；接口文档当前未明确该字段。"),
        ("工位预约", "/system/reservation/detailList", "可复用", "管理端需要查询全部记录；不能强制 userId 为当前用户。"),
        ("塔服订单", "/system/order/detailList", "需扩展", "统一 orderTime/createTime、paymentProofUrl/paymentProofPath 字段命名。"),
        ("塔服收款", "PUT /system/order", "可复用", "理论上可用通用修改，但专用 approve 更适合状态机校验、权限和审计。"),
        ("学习资料上传", "/common/upload", "需扩展", "上传目录白名单需加入资料目录，并补资料元数据 CRUD 与下载权限。"),
        ("招新导出", "标准 /export", "可复用", "若按若依惯例建议 POST 文件流；必须接受与列表相同筛选条件。"),
    ]
    add_table(doc, ["页面", "现有接口", "判断", "需要后端确认/扩展"], reuse_rows, [1500, 2600, 1100, 4160], font_size=8, status_col=2)

    add_heading(doc, "5. 建议的缺失接口契约", 1)
    contract_rows = [
        ("分页列表", "pageNum、pageSize、筛选字段", "TableDataInfo：total、rows、code、msg"),
        ("成员届次", "可选 status/year", "currentCohort、cohorts[]；每届含 id/name/isCurrent"),
        ("留任", "userId、sourceCohortId、targetCohortId、retain=true", "原届状态、目标届成员、是否新建届次；重复请求不重复建数据"),
        ("招新列表", "round、departmentId、status、pageNum、pageSize", "申请人、两志愿及各自状态、投递时间、操作权限"),
        ("面评", "applicationId、roundId、departmentId、content", "evaluationId、evaluator、created/updatedTime；只能修改本人面评"),
        ("二面录用", "applicationId、volunteerNo/departmentId、decision(PASS/OUT)、notice", "最终状态、通知状态；两个志愿均通过时默认第一志愿"),
        ("报名聚合", "activityId/type、pageNum、pageSize", "total、quota、rows；rows 包含页面展示的学生与备注字段"),
        ("资料上传", "multipart file + category + visibility", "materialId、fileName、size、uploader、uploadTime、downloadUrl"),
        ("收款确认", "orderId", "paymentStatus、confirmedBy、confirmedAt；重复确认应幂等"),
        ("文件下载", "资源 id 或 filename", "blob/file stream；Content-Disposition 提供文件名"),
    ]
    add_table(doc, ["场景", "建议请求", "建议返回/规则"], contract_rows, [1500, 3300, 4560], font_size=8)

    add_heading(doc, "6. 权限、数据范围与事务要求", 1)
    security_rows = [
        ("成员删除/重置/导入", "仅 CEO", "后端权限注解 + 业务角色二次校验；不能只靠按钮隐藏。"),
        ("成员只读", "其他管理层", "可查看成员管理，但无删除、重置、导入、留任修改权限。"),
        ("招新管理", "管理层", "除 CEO 外严格限制为本部门候选人；列表、详情、面评、评定、导出均需同一数据范围。"),
        ("招新二面面评", "经理层", "只允许编辑本部门二面面评，不显示工位预约与塔服订购菜单。"),
        ("学习资料", "管理端角色/上传者", "上传、删除、修改、下载分别校验权限与可见范围。"),
        ("留任换届", "CEO + 定时任务", "留任同步、新建届次、源状态更新必须在同一事务中并支持幂等。"),
        ("面试最终结果", "管理层本部门", "Pass/Out、通知生成、双志愿默认规则应原子提交并留审计日志。"),
        ("后台 CRUD", "按资源分权", "立即恢复 @PreAuthorize；敏感资源补部门数据权限与操作日志。"),
    ]
    add_table(doc, ["能力", "角色/范围", "后端要求"], security_rows, [2100, 2200, 5060], font_size=8.3)

    add_heading(doc, "7. 建议后端确认顺序", 1)
    priority_rows = [
        ("P0", "鉴权与数据范围", "恢复后台 CRUD 鉴权；确认 CEO、管理层、经理层权限及部门隔离。"),
        ("P0", "成员管理", "届次、留任、删除、重置密码及 8 月 1 日换届任务。"),
        ("P0", "招新管理", "管理端列表/详情/面评/评定/通知/双志愿规则/导出。"),
        ("P0", "塔服收款", "确认收款专用接口或明确通用 PUT 的状态机契约。"),
        ("P1", "活动名单", "确认是否复用 activity/signup，补足人数、名额和报名人字段。"),
        ("P1", "学习资料", "资料 CRUD、上传目录、下载、可见范围。"),
        ("P1", "字段统一", "成员、图书、订单、预约的页面字段与接口字段逐项对齐。"),
        ("P2", "配置化", "招新轮次、部门、枚举字典与控制台统计。"),
    ]
    add_table(doc, ["优先级", "主题", "需确认事项"], priority_rows, [1000, 2100, 6260], font_size=8.5)

    add_heading(doc, "8. 前端当前接入状态", 1)
    add_table(doc, ["状态", "说明"], [
        ("已写页面与 API 封装", "成员管理、招新管理、宣讲会报名、精英分享会、工位预约记录、图书借阅记录、塔服订购、学习资料等。"),
        ("已按后端文档定义", "登录、用户信息、动态路由、退出、验证码、成员列表、部分记录列表与面试结果录入。"),
        ("暂时使用 Mock", "后端尚未上线，开发环境 VITE_USE_MOCK=true；缺失接口的交互仅为前端模拟。"),
        ("联调前必须完成", "后端确认路径、方法、参数、响应字段、文件流处理、权限码、角色与部门数据范围。"),
    ], [1900, 7460], font_size=9)

    add_note(doc, "最终说明", "本文件用于接口比对，不代表后端接口已经实现。以根目录《接口文档.md》为“已明确”依据；前端 src/api 只代表页面期望的调用契约。", LIGHT_BLUE)

    doc.core_properties.title = "Quanta 后台管理系统接口对照说明"
    doc.core_properties.subject = "后端现有接口与前端缺失接口比对"
    doc.core_properties.author = "Quanta 项目组"
    doc.save(OUTPUT)
    print(str(OUTPUT).encode("unicode_escape").decode("ascii"))


if __name__ == "__main__":
    main()
